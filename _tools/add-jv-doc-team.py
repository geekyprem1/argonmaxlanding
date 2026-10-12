from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import io, re

root=Path(__file__).resolve().parent.parent
path=root/'AroganMax AI JV Doc.docx'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
P='http://schemas.openxmlformats.org/package/2006/relationships'
C='http://schemas.openxmlformats.org/package/2006/content-types'
ns={'w':W,'r':R,'a':A,'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
with ZipFile(path) as z:
    original={i.filename:z.read(i.filename) for i in z.infolist()}
    infos=z.infolist()
docxml=E.fromstring(original['word/document.xml'])
assert 'Meet the team behind ArgonMAX AI' not in ''.join(docxml.xpath('//w:t/text()',namespaces=ns)), 'Team section already present'
frag=Document()
table=frag.add_table(rows=4, cols=3)
table.alignment=WD_TABLE_ALIGNMENT.CENTER
table.autofit=False
for col,width in zip(table.columns,[3.11,.28,3.11]): col.width=Inches(width)
for row in table.rows:
    pr=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(pr)
    for cell,width in zip(row.cells,[3.11,.28,3.11]):
        cell.width=Inches(width)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        tcPr=cell._tc.get_or_add_tcPr()
        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'0B1D35'); tcPr.append(shade)
        margins=OxmlElement('w:tcMar')
        for edge,amount in [('top',200),('bottom',200),('left',220),('right',220)]:
            el=OxmlElement('w:'+edge); el.set(qn('w:w'),str(amount)); el.set(qn('w:type'),'dxa'); margins.append(el)
        tcPr.append(margins)
props=table._tbl.tblPr
borders=OxmlElement('w:tblBorders')
for edge in ['top','left','bottom','right','insideH','insideV']:
    el=OxmlElement('w:'+edge); el.set(qn('w:val'),'nil'); borders.append(el)
props.append(borders)

# All section formatting is direct, avoiding changes to existing document styles.
def paragraph(cell,text='',size=10.5,color='C5D2E3',bold=False,after=8,before=0,first=False):
    p=cell.paragraphs[0] if first else cell.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pf=p.paragraph_format; pf.space_after=Pt(after); pf.space_before=Pt(before); pf.line_spacing=1.12; pf.keep_with_next=False
    r=p.add_run(text); r.font.name='Inter'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color); r.bold=bold
    return p

top=table.cell(0,0).merge(table.cell(0,2))
paragraph(top,'MEET THE FOUNDERS',10,'7BDCF5',True,after=12,first=True)
paragraph(top,'Meet the team behind ArgonMAX AI',23,'FFFFFF',True,after=10)
intro=table.cell(1,0).merge(table.cell(1,2))
paragraph(intro,'Product strategy meets AI engineering. Two dedicated minds, one mission: helping creators and businesses bring their ideas to life.',10.5,after=14,first=True)
profiles=[
 ('Pradip','Product Strategist','pradip.jpg',1216,1038,'Leads the product vision, strategy and customer-focused innovation behind ArgonMAX AI.','Product Planning · Marketing Strategy · Funnel Optimization','https://www.facebook.com/pradippanditpk/'),
 ('Prem','AI System Engineer','prem.jpg',1600,900,'Builds the AI systems powering ArgonMAX AI, with advanced automation and intelligent workflows.','AI Development · Machine Learning · Software Architecture','https://www.facebook.com/Prempandit')
]
for col,person in zip([0,2],profiles):
    name,role,photo,width,height,description,expertise,url=person
    cell=table.cell(2,col)
    cell._tc.get_or_add_tcPr().find(qn('w:shd')).set(qn('w:fill'),'142B45')
    p=paragraph(cell,after=13,first=True)
    shape=p.add_run().add_picture(str(root/'assets/jv/founders'/photo),width=Inches(1.5),height=Inches(1.5))
    shape._inline.docPr.set('descr',name+' '+role)
    blipFill=shape._inline.xpath('.//pic:blipFill')[0]
    crop=OxmlElement('a:srcRect')
    side=round((width-height)/width*50000)
    crop.set('l',str(side));crop.set('r',str(side));crop.set('t','0');crop.set('b','0')
    blipFill.insert(1,crop)
    geom=shape._inline.xpath('.//a:prstGeom')[0]; geom.set('prst','ellipse')
    spPr=shape._inline.xpath('.//pic:spPr')[0]
    ln=OxmlElement('a:ln'); ln.set('w','25400'); solid=OxmlElement('a:solidFill'); color=OxmlElement('a:srgbClr');color.set('val','55CAEF');solid.append(color);ln.append(solid);spPr.append(ln)
    paragraph(cell,name,18,'FFFFFF',True,after=5)
    paragraph(cell,role,11,'7BDCF5',True,after=14)
    paragraph(cell,description,10.5,after=16)
    paragraph(cell,'EXPERTISE',8.5,'92CDE5',True,after=6)
    paragraph(cell,expertise,9.5,after=16)
    p=paragraph(cell,after=6)
    link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),frag.part.relate_to(url,RT.HYPERLINK,is_external=True));link.set(qn('w:history'),'1')
    run=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Inter');font.set(qn('w:hAnsi'),'Inter');rp.append(font)
    colr=OxmlElement('w:color');colr.set(qn('w:val'),'7BDCF5');rp.append(colr)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'21');rp.append(size)
    bold=OxmlElement('w:b');rp.append(bold);u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u)
    run.append(rp);t=OxmlElement('w:t');t.text='Connect with '+name+' on Facebook';run.append(t);link.append(run);p._p.append(link)
quote=table.cell(3,0).merge(table.cell(3,2))
paragraph(quote,'Together, we are building a simpler way to create with AI.',11,after=4,before=10,first=True)

buf=io.BytesIO();frag.save(buf)
with ZipFile(buf) as z: parts={n:z.read(n) for n in z.namelist()}
rels=E.fromstring(original['word/_rels/document.xml.rels'])
fragRels=E.fromstring(parts['word/_rels/document.xml.rels'])
fragment=E.fromstring(parts['word/document.xml']).find('w:body/w:tbl',ns)
ct=E.fromstring(original['[Content_Types].xml'])
usedIds={r.get('Id') for r in rels}
newParts={}
for attr in ['{'+R+'}id','{'+R+'}embed']:
    for node in fragment.iter():
        old=node.get(attr)
        if not old: continue
        relation=next(x for x in fragRels if x.get('Id')==old)
        newId='rIdTeam'+str(len(usedIds)+1)
        while newId in usedIds: newId+='x'
        usedIds.add(newId)
        newRel=deepcopy(relation);newRel.set('Id',newId)
        if relation.get('Type')==RT.IMAGE:
            src='word/'+relation.get('Target')
            filename='argonmax-team-'+('pradip' if 'pradip' in node.getparent().getparent().getparent().getparent().getparent().getparent().tag else '')
            target='media/argonmax-team-'+str(len(newParts)+1)+'.jpg'
            while 'word/'+target in original: target=target.replace('.jpg','x.jpg')
            newParts['word/'+target]=parts[src]
            newRel.set('Target',target)
        rels.append(newRel);node.set(attr,newId)
# Unique drawing IDs across the complete existing document.
ids=[int(x.get('id')) for x in docxml.findall('.//wp:docPr',ns)]
for i,draw in enumerate(fragment.findall('.//wp:docPr',ns),start=max(ids,default=0)+1): draw.set('id',str(i))
if not any(x.get('Extension','').lower()=='jpg' for x in ct):
    el=E.SubElement(ct,'{'+C+'}Default');el.set('Extension','jpg');el.set('ContentType','image/jpeg')
body=docxml.find('w:body',ns)
page=E.Element('{'+W+'}p');pr=E.SubElement(page,'{'+W+'}pPr');E.SubElement(pr,'{'+W+'}pageBreakBefore');spacing=E.SubElement(pr,'{'+W+'}spacing');spacing.set('{'+W+'}before','0');spacing.set('{'+W+'}after','0');spacing.set('{'+W+'}line','20');spacing.set('{'+W+'}lineRule','exact')
body.insert(len(body)-1,page);body.insert(len(body)-1,fragment)
# Word requires a trailing paragraph after a table.
end=E.Element('{'+W+'}p');pr=E.SubElement(end,'{'+W+'}pPr');spacing=E.SubElement(pr,'{'+W+'}spacing');spacing.set('{'+W+'}after','0');spacing.set('{'+W+'}line','20');spacing.set('{'+W+'}lineRule','exact');body.insert(len(body)-1,end)
serialize=lambda element:E.tostring(element,xml_declaration=True,encoding='UTF-8',standalone=True)
changes={'word/document.xml':serialize(docxml),'word/_rels/document.xml.rels':serialize(rels),'[Content_Types].xml':serialize(ct)}
tmp=path.with_suffix('.team.tmp.docx')
with ZipFile(tmp,'w',ZIP_DEFLATED) as out:
    for info in infos: out.writestr(info,changes.get(info.filename,original[info.filename]))
    for n,data in newParts.items(): out.writestr(n,data)
tmp.replace(path)
with ZipFile(path) as z:
    assert z.testzip() is None
    actual=E.fromstring(z.read('word/document.xml'))
    actualBody=actual.find('w:body',ns)
    oldBody=E.fromstring(original['word/document.xml']).find('w:body',ns)
    assert len(actualBody)==len(oldBody)+3
    for before,after in zip(list(oldBody)[:-1],list(actualBody)[:len(oldBody)-1]):
        assert E.tostring(before,method='c14n')==E.tostring(after,method='c14n'), 'Original document content changed'
    urls=[x.get('Target') for x in E.fromstring(z.read('word/_rels/document.xml.rels')) if x.get('Type')==RT.HYPERLINK]
    for person in profiles: assert person[-1] in urls
    assert len(newParts)==2
print('Updated',path.name,'with 2 embedded photos and 2 clickable Facebook links. Original body preserved.')
